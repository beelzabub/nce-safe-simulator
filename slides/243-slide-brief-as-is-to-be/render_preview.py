#!/usr/bin/env python3
"""Rasterize a .pptx to per-slide PNGs with Pillow — a LibreOffice-free preview.

Reads each shape's real geometry, fill, and per-run colors straight from the
pptx, so it faithfully reveals layout/alignment/color bugs (it does NOT re-derive
from the builder). Fonts are substituted (Gill Sans MT -> DejaVu Sans,
Consolas -> DejaVu Sans Mono), so previews are layout-faithful, not font-faithful.

    python3 render_preview.py [deck.pptx] [out_dir]
"""
import io
import os
import sys

from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image, ImageDraw, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    HERE, "dist", "NCE-Safe-Simulator-SDLC-Modernization.pptx")
OUTDIR = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "preview")

DPI = 150
PXPI = DPI / 914400.0          # px per EMU
PTPX = DPI / 72.0              # px per point

FONTS = {
    (False, False): "/usr/share/fonts/dejavu-sans-fonts/DejaVuSans.ttf",
    (False, True): "/usr/share/fonts/dejavu-sans-fonts/DejaVuSans-Bold.ttf",
    ("mono", False): "/usr/share/fonts/dejavu-sans-mono-fonts/DejaVuSansMono.ttf",
    ("mono", True): "/usr/share/fonts/dejavu-sans-mono-fonts/DejaVuSansMono-Bold.ttf",
    ("it", False): "/usr/share/fonts/dejavu-sans-fonts/DejaVuSans-Oblique.ttf",
}
_font_cache = {}


def font_for(name, bold, italic, px):
    mono = name and "Mono" in name or name in ("Consolas", "Courier New")
    if mono:
        key = ("mono", bool(bold))
    elif italic:
        key = ("it", bool(bold))
    else:
        key = (False, bool(bold))
    path = FONTS.get(key, FONTS[(False, False)])
    ck = (path, px)
    if ck not in _font_cache:
        _font_cache[ck] = ImageFont.truetype(path, max(6, int(px)))
    return _font_cache[ck]


def emu(v):
    return int((v or 0) * PXPI)


def run_color(run, default=(20, 30, 40)):
    try:
        c = run.font.color
        if c and c.type is not None and c.rgb is not None:
            return (c.rgb[0], c.rgb[1], c.rgb[2])
    except Exception:
        pass
    return default


def solid_fill(shape):
    try:
        if shape.fill.type == 1:  # MSO_FILL.SOLID
            rgb = shape.fill.fore_color.rgb
            return (rgb[0], rgb[1], rgb[2])
    except Exception:
        pass
    return None


def wrap_runs(runs, box_w, word_wrap):
    """runs: list of (text, font, color). Returns lines: list of list of
    (text, font, color, width). Greedy word wrap when word_wrap and box_w."""
    lines = [[]]
    x = 0
    for text, font, color in runs:
        if not text:
            continue
        # split keeping spaces attached to preceding word
        tokens = []
        cur = ""
        for ch in text:
            cur += ch
            if ch == " ":
                tokens.append(cur)
                cur = ""
        if cur:
            tokens.append(cur)
        for tok in tokens:
            w = font.getlength(tok)
            if word_wrap and box_w and x + w > box_w and lines[-1]:
                lines.append([])
                x = 0
                tok_ls = tok.lstrip()
                w = font.getlength(tok_ls)
                tok = tok_ls
            lines[-1].append((tok, font, color, w))
            x += w
    return lines


def draw_text_frame(draw, tf, left, top, width, height):
    pad_l = emu(getattr(tf, "margin_left", 0) or 0)
    pad_r = emu(getattr(tf, "margin_right", 0) or 0)
    pad_t = emu(getattr(tf, "margin_top", 0) or 0)
    pad_b = emu(getattr(tf, "margin_bottom", 0) or 0)
    box_w = width - pad_l - pad_r
    word_wrap = bool(getattr(tf, "word_wrap", True))
    # build all display lines with metrics
    disp = []          # (line_runs, line_h, align)
    for para in tf.paragraphs:
        align = para.alignment or PP_ALIGN.LEFT
        runs = []
        maxpx = 12
        for r in para.runs:
            sz = r.font.size.pt if r.font.size else 14
            px = sz * PTPX
            maxpx = max(maxpx, px)
            f = font_for(r.font.name, r.font.bold, r.font.italic, px)
            runs.append((r.text, f, run_color(r)))
        if not runs:
            disp.append(([], int(maxpx * 1.3), align))
            continue
        for ln in wrap_runs(runs, box_w, word_wrap):
            disp.append((ln, int(maxpx * 1.32), align))
    total_h = sum(h for _, h, _ in disp)
    anchor = getattr(tf, "vertical_anchor", None)
    if anchor == MSO_ANCHOR.MIDDLE:
        y = top + pad_t + max(0, (height - pad_t - pad_b - total_h) // 2)
    elif anchor == MSO_ANCHOR.BOTTOM:
        y = top + height - pad_b - total_h
    else:
        y = top + pad_t
    for ln, lh, align in disp:
        line_w = sum(w for _, _, _, w in ln)
        if align == PP_ALIGN.CENTER:
            x = left + pad_l + max(0, (box_w - line_w) // 2)
        elif align == PP_ALIGN.RIGHT:
            x = left + pad_l + max(0, box_w - line_w)
        else:
            x = left + pad_l
        for text, f, color, w in ln:
            draw.text((x, y), text, font=f, fill=color)
            x += w
        y += lh


def draw_table(draw, shape):
    tbl = shape.table
    x0, y0 = emu(shape.left), emu(shape.top)
    col_w = [emu(c.width) for c in tbl.columns]
    row_h = [emu(r.height) for r in tbl.rows]
    y = y0
    for ri, row in enumerate(tbl.rows):
        x = x0
        for ci in range(len(col_w)):
            cell = tbl.cell(ri, ci)
            w, h = col_w[ci], row_h[ri]
            fill = solid_fill(cell)
            if fill:
                draw.rectangle([x, y, x + w, y + h], fill=fill)
            draw.rectangle([x, y, x + w, y + h], outline=(210, 216, 222))
            draw_text_frame(draw, cell.text_frame, x, y, w, h)
            x += w
        y += row_h[ri]


def render_slide(slide, idx):
    W = emu(slide.part.package.presentation_part.presentation.slide_width)
    H = emu(slide.part.package.presentation_part.presentation.slide_height)
    img = Image.new("RGB", (W, H), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    for shape in slide.shapes:
        try:
            if shape.shape_type == 13:  # PICTURE
                blob = shape.image.blob
                pic = Image.open(io.BytesIO(blob)).convert("RGBA")
                pic = pic.resize((max(1, emu(shape.width)), max(1, emu(shape.height))))
                img.paste(pic, (emu(shape.left), emu(shape.top)), pic)
                continue
            if shape.has_table:
                draw_table(draw, shape)
                continue
            if hasattr(shape, "begin_x") and hasattr(shape, "end_x"):  # connector
                import math
                x1, y1 = emu(shape.begin_x), emu(shape.begin_y)
                x2, y2 = emu(shape.end_x), emu(shape.end_y)
                col = (245, 196, 76)
                try:
                    c = shape.line.color
                    if c and c.type is not None and c.rgb is not None:
                        col = (c.rgb[0], c.rgb[1], c.rgb[2])
                except Exception:
                    pass
                draw.line([(x1, y1), (x2, y2)], fill=col, width=4)
                ang = math.atan2(y2 - y1, x2 - x1)
                for da in (2.62, -2.62):  # ~150 degrees, arrowhead at the end
                    draw.line([(x2, y2), (x2 + 16 * math.cos(ang + da),
                                         y2 + 16 * math.sin(ang + da))], fill=col, width=4)
                continue
            fill = solid_fill(shape)
            if fill is not None and shape.width and shape.height:
                box = [emu(shape.left), emu(shape.top),
                       emu(shape.left) + emu(shape.width), emu(shape.top) + emu(shape.height)]
                is_round = "ROUNDED" in str(getattr(shape, "auto_shape_type", ""))
                if is_round:
                    draw.rounded_rectangle(box, radius=int(0.08 * emu(shape.height)), fill=fill)
                else:
                    draw.rectangle(box, fill=fill)
            if shape.has_text_frame:
                draw_text_frame(draw, shape.text_frame, emu(shape.left), emu(shape.top),
                                emu(shape.width), emu(shape.height))
        except Exception as e:
            print(f"  slide {idx} shape skip: {e}")
    return img


def main():
    os.makedirs(OUTDIR, exist_ok=True)
    prs = Presentation(DECK)
    for i, slide in enumerate(prs.slides, 1):
        img = render_slide(slide, i)
        out = os.path.join(OUTDIR, f"slide-{i:02d}.png")
        img.save(out)
        print(f"wrote {out}")


if __name__ == "__main__":
    main()
