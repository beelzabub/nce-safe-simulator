"""
Generate deck/assets/template.pptx — the status deck's template of record.

The deck used to fetch its template from S3 (the SAIC branding file). Instead the
template is now committed in the repo (issue #210): this script strips every
content slide out of a source .pptx, leaving only the slide masters, layouts, and
theme — the styling the deck is built on. build_deck.py creates the cover from
the Cover 1 layout, so the template needs no seed slides.

Usage:
  python3 deck/make_template.py [--source deck/dist/<built>.pptx] [--out deck/assets/template.pptx]

With no --source it strips the newest built deck under deck/dist/. To rebuild the
template from the original SAIC file, first fetch it once with
`build_deck.py --template-s3-path s3://...`, build a deck, then run this.
"""
import argparse
import glob
import os

from pptx import Presentation
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))


def strip_to_template(src, out):
    """Remove every slide from src, keeping masters/layouts/theme, and save to out."""
    prs = Presentation(src)
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.get(qn("r:id")))
        sld_id_lst.remove(sld_id)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    return Presentation(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", help="source pptx (default: newest deck/dist/*.pptx)")
    ap.add_argument("--out", default=os.path.join(HERE, "assets", "template.pptx"))
    args = ap.parse_args()

    src = args.source
    if not src:
        cands = sorted(glob.glob(os.path.join(HERE, "dist", "*.pptx")), key=os.path.getmtime)
        if not cands:
            raise SystemExit("no deck/dist/*.pptx found — build a deck first, or pass --source")
        src = cands[-1]

    tmpl = strip_to_template(src, args.out)
    print(f"Wrote {args.out} from {src}")
    print(f"  {len(tmpl.slides)} content slides, {len(tmpl.slide_masters)} masters, "
          f"{os.path.getsize(args.out) / 1e6:.1f} MB")


if __name__ == "__main__":
    main()
